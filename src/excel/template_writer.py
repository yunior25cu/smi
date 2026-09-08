"""Escribe valores de indicadores en una copia de la plantilla Excel de una institución.

Revisión 2 del prompt (cambio de lógica): el entregable real es la hoja
``4315 Utilizacion `` (con espacio final, no ``DJ``), y dentro de esa hoja
cada indicador tiene un par de columnas contiguas no simétricas:

- **Columna PRECIO** (ej. ``G22``): valor administrativo ya cargado y
  correcto. Nunca se escribe, nunca se toca.
- **Columna ORDENES/TICKET** (ej. ``H22``, inmediatamente a la derecha):
  hoy tiene, como placeholder, el texto de ``PLASQLDESC`` del indicador que
  le corresponde. Es la única celda que este módulo escribe, reemplazando
  ese texto por el valor numérico resultante de ``PLASQLSENT``.

Como ``write_values`` solo escribe las celdas que vienen en ``values``
(nunca las que no están en el mapa), la columna PRECIO queda preservada por
construcción con solo no incluirla nunca en ese diccionario -no hace falta
lógica extra para "no tocarla". Lo que sí es nuevo es la verificación
post-escritura: después de escribir, se recorren todas las columnas
ORDENES/TICKET de la hoja (detectadas por encabezado, no por posición fija)
y se reportan las que todavía tengan texto en vez de un número -señal de un
indicador que se quedó sin mapear o de una celda mal apuntada en
``indicator_map.csv``.

Nunca modifica el archivo plantilla original (decisión 5): primero se copia
el archivo a la ruta de salida, y recién ahí se abre la copia con openpyxl
para escribir encima.

No hace falta la contraseña de protección de hoja para escribir: se verificó
en la Fase 0, contra la plantilla real convertida, que openpyxl escribe
valores en una hoja protegida sin pedirla, y que el archivo resultante
conserva la protección y las celdas combinadas intactas.
"""

from __future__ import annotations

import logging
import shutil
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)

# Confirmado con el usuario: las celdas ORDENES/TICKET representan una
# cantidad (órdenes/tickets), no un monto, así que van sin decimales.
DEFAULT_NUMBER_FORMAT = "#,##0"

# Encabezado literal que identifica una columna ORDENES/TICKET en la fila de
# encabezados de la hoja (ver decisión 4 del prompt de cambio de lógica).
ORDENES_TICKET_HEADER = "ORDENES/TICKET"
PRECIO_HEADER_PREFIX = "PRECIO"


class TemplateWriterError(ValueError):
    """Error al escribir la plantilla (falta el archivo, falta la hoja, etc.)."""


@dataclass(frozen=True)
class WriteResult:
    output_path: Path
    # (celda, texto) por cada columna ORDENES/TICKET que, después de
    # escribir, sigue conteniendo texto en vez de un número.
    celdas_con_texto_residual: tuple[tuple[str, str], ...]


def write_values(
    template_path: str | Path,
    output_path: str | Path,
    hoja: str,
    values: dict[str, float | None],
    *,
    number_format: str = DEFAULT_NUMBER_FORMAT,
) -> WriteResult:
    """Copia ``template_path`` a ``output_path`` y escribe ``values`` en ``hoja``.

    ``values`` es {celda: valor}, ej. {"H22": 346.0, "H23": None}. Un valor
    None se escribe tal cual (celda vacía): decidir qué hacer con los
    indicadores sin resultado es responsabilidad del log de control, no de
    este módulo. Cada celda escrita recibe ``number_format`` (por defecto
    ``#,##0``, confirmado con el usuario) -la celda hoy es texto plano, así
    que no hay un formato numérico previo del que partir.

    Nunca escribe nada fuera de ``values``: si ``values`` no incluye una
    celda PRECIO, esa celda queda intacta, byte a byte.

    Devuelve un WriteResult con la ruta de salida y las celdas ORDENES/TICKET
    que, tras escribir, todavía tienen texto en vez de un número (indicador
    sin mapear o celda mal apuntada). Si ``output_path`` ya existe, se
    sobreescribe (volver a correr una exportación para la misma
    institución/mes es normal, no un error), pero se deja constancia en el
    log.
    """
    template_path = Path(template_path)
    output_path = Path(output_path)

    if not template_path.exists():
        raise TemplateWriterError(f"no existe la plantilla {template_path}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        logger.info("%s ya existía, se sobreescribe", output_path)
    shutil.copy2(template_path, output_path)

    wb = openpyxl.load_workbook(output_path)
    if hoja not in wb.sheetnames:
        raise TemplateWriterError(
            f"la plantilla {template_path} no tiene una hoja llamada {hoja!r} "
            f"(hojas disponibles: {wb.sheetnames})"
        )
    sheet = wb[hoja]

    for celda, valor in values.items():
        cell = sheet[celda]
        cell.value = valor
        cell.number_format = number_format

    leftovers = tuple(find_leftover_placeholder_cells(sheet))
    wb.save(output_path)
    return WriteResult(output_path=output_path, celdas_con_texto_residual=leftovers)


def find_ordenes_ticket_columns(sheet: Worksheet, search_rows: int = 40) -> tuple[int, list[int]]:
    """Ubica la fila de encabezados y las columnas ORDENES/TICKET de la hoja.

    Detecta las columnas por su encabezado literal, no por una posición fija
    (evita hardcodear letras de columna, que además difieren de la plantilla
    vieja DJ). Excluye la columna "TOTAL" -tiene el mismo encabezado
    "ORDENES/TICKET" pero es una fórmula de suma, no la celda de un
    indicador individual-: una columna cuenta solo si tiene una columna
    "PRECIO N" inmediatamente a la izquierda, en la fila de arriba.
    """
    for row in range(1, search_rows + 1):
        cols = []
        for col in range(1, sheet.max_column + 1):
            if sheet.cell(row, col).value != ORDENES_TICKET_HEADER:
                continue
            precio_header = sheet.cell(row - 1, col - 1).value
            if isinstance(precio_header, str) and precio_header.strip().upper().startswith(PRECIO_HEADER_PREFIX):
                cols.append(col)
        if cols:
            return row, cols
    raise TemplateWriterError(
        f"no se encontraron columnas {ORDENES_TICKET_HEADER!r} en las primeras {search_rows} filas de la hoja"
    )


def find_leftover_placeholder_cells(sheet: Worksheet) -> list[tuple[str, str]]:
    """Recorre las columnas ORDENES/TICKET de la hoja y devuelve (celda, texto)
    por cada una que todavía tenga texto (no un número).

    No filtra por "fila con concepto en columna A": esta hoja tiene celdas de
    concepto combinadas verticalmente (una sola celda de texto cubriendo
    varias filas), así que las filas dentro de un combinado devuelven
    ``None`` en columna A aunque tengan su propio valor real en las columnas
    ORDENES/TICKET -filtrar por columna A perdería esas filas. Cualquier fila
    sin datos en ninguna columna ORDENES/TICKET simplemente no aporta nada al
    resultado, así que no hace falta el filtro para que esto sea correcto.
    """
    header_row, columns = find_ordenes_ticket_columns(sheet)

    leftovers: list[tuple[str, str]] = []
    for row in range(header_row + 1, sheet.max_row + 1):
        for col in columns:
            cell = sheet.cell(row, col)
            if isinstance(cell.value, str) and cell.value.strip() != "":
                leftovers.append((cell.coordinate, cell.value.strip()))
    return leftovers
