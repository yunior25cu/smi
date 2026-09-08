"""Utilidad de un solo uso: propone filas de config/indicator_map.csv
cruzando el texto placeholder que hoy ocupa cada celda ORDENES/TICKET de la
hoja "4315 Utilizacion " contra PLACONSU.PLASQLDESC.

No es un mapeo final automático (decisión 2 del prompt de cambio de lógica):
un texto duplicado en el catálogo (varios plasqlid con el mismo PLASQLDESC)
se marca como **ambiguo**, y un texto sin match se marca como **huérfano**.
Ninguno de los dos entra al CSV propuesto -ambos quedan en un reporte aparte
para revisión humana. No se usa en la corrida mensual regular (no forma
parte de run_export.py).

Uso:
    python -m tools.build_indicator_map --template "templates/SMI/4315 Utilizacion.xlsx" \
        --institucion SMI --out-csv indicator_map_propuesto.csv \
        --out-report indicator_map_revision.txt
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

import openpyxl

from src.db.catalog_repository import CatalogEntry
from src.excel.template_writer import find_leftover_placeholder_cells

DEFAULT_HOJA = "4315 Utilizacion "


@dataclass(frozen=True)
class ProposedMapping:
    plasqlid: int
    celda: str
    texto: str


@dataclass(frozen=True)
class AmbiguousCell:
    celda: str
    texto: str
    candidatos: tuple[int, ...]


@dataclass(frozen=True)
class OrphanCell:
    celda: str
    texto: str


@dataclass(frozen=True)
class DuplicatePlasqlid:
    """Un mismo plasqlid matcheó, cada uno de forma no ambigua, a más de una
    celda distinta -típicamente porque el mismo texto placeholder aparece
    repetido en dos columnas de la plantilla (ej. bloque NO FONASA y FONASA)
    apuntando al mismo indicador del catálogo. indicator_map.py rechazaría
    el CSV final si esto llegara sin resolver, así que se separa acá para
    revisión humana en vez de dejar pasar una sola de las celdas al azar."""

    plasqlid: int
    celdas: tuple[str, ...]


@dataclass(frozen=True)
class MapBuildResult:
    confirmados: tuple[ProposedMapping, ...]
    ambiguos: tuple[AmbiguousCell, ...]
    huerfanos: tuple[OrphanCell, ...]
    duplicados: tuple[DuplicatePlasqlid, ...] = ()


def build_proposal(
    placeholders: Sequence[tuple[str, str]], catalog: Sequence[CatalogEntry]
) -> MapBuildResult:
    """Cruza (celda, texto) contra el catálogo por PLASQLDESC primero y, si no
    hay match, por PLASQLNOM como respaldo (sin mayúsculas/minúsculas ni
    espacios en los extremos). El respaldo por PLASQLNOM importa: se
    verificó contra el catálogo real que algunas filas tienen PLASQLDESC
    vacío, y el placeholder de esa celda quedó armado con PLASQLNOM en su
    lugar (ej. plasqlid 640, "LARINGOSCOPIA ORDENES NF P1")."""
    by_desc: dict[str, list[int]] = {}
    by_nombre: dict[str, list[int]] = {}
    for entry in catalog:
        by_desc.setdefault(entry.descripcion.strip().upper(), []).append(entry.plasqlid)
        by_nombre.setdefault(entry.nombre.strip().upper(), []).append(entry.plasqlid)

    confirmados: list[ProposedMapping] = []
    ambiguos: list[AmbiguousCell] = []
    huerfanos: list[OrphanCell] = []

    for celda, texto in placeholders:
        key = texto.strip().upper()
        candidatos = sorted(set(by_desc.get(key, [])))
        if not candidatos:
            candidatos = sorted(set(by_nombre.get(key, [])))
        if len(candidatos) == 1:
            confirmados.append(ProposedMapping(plasqlid=candidatos[0], celda=celda, texto=texto))
        elif len(candidatos) > 1:
            ambiguos.append(AmbiguousCell(celda=celda, texto=texto, candidatos=tuple(candidatos)))
        else:
            huerfanos.append(OrphanCell(celda=celda, texto=texto))

    # Segunda pasada: un plasqlid que matcheó (sin ambigüedad en ese momento)
    # a más de una celda distinta es, en conjunto, un conflicto -se saca de
    # "confirmados" entero, no se elige una celda al azar.
    celdas_por_id: dict[int, set[str]] = {}
    for m in confirmados:
        celdas_por_id.setdefault(m.plasqlid, set()).add(m.celda)

    ids_en_conflicto = {pid for pid, celdas in celdas_por_id.items() if len(celdas) > 1}
    duplicados = tuple(
        sorted(
            (
                DuplicatePlasqlid(plasqlid=pid, celdas=tuple(sorted(celdas_por_id[pid])))
                for pid in ids_en_conflicto
            ),
            key=lambda d: d.plasqlid,
        )
    )
    confirmados = [m for m in confirmados if m.plasqlid not in ids_en_conflicto]

    return MapBuildResult(
        confirmados=tuple(sorted(confirmados, key=lambda m: m.celda)),
        ambiguos=tuple(sorted(ambiguos, key=lambda a: a.celda)),
        huerfanos=tuple(sorted(huerfanos, key=lambda o: o.celda)),
        duplicados=duplicados,
    )


def build_proposal_from_workbook(
    template_path: str | Path, hoja: str, catalog: Sequence[CatalogEntry]
) -> MapBuildResult:
    wb = openpyxl.load_workbook(template_path)
    if hoja not in wb.sheetnames:
        raise ValueError(f"{template_path}: no tiene una hoja llamada {hoja!r} (hojas: {wb.sheetnames})")
    placeholders = find_leftover_placeholder_cells(wb[hoja])
    return build_proposal(placeholders, catalog)


def write_indicator_map_csv(
    result: MapBuildResult, institucion_id: str, hoja: str, path: str | Path
) -> Path:
    """Escribe solo los mapeos CONFIRMADOS (sin ambigüedad ni huérfanos) en
    formato indicator_map.csv, listo para revisar y mergear a mano en
    config/indicator_map.csv -esta función nunca escribe ahí directamente."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        f.write("plasqlid,institucion_id,hoja,celda\n")
        for m in result.confirmados:
            f.write(f"{m.plasqlid},{institucion_id},{hoja},{m.celda}\n")
    return path


def write_review_report(result: MapBuildResult, path: str | Path) -> Path:
    """Reporte legible por humano de lo que necesita revisión manual."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        f"=== Propuesta de mapeo: {len(result.confirmados)} confirmados, "
        f"{len(result.ambiguos)} ambiguos, {len(result.huerfanos)} huérfanos, "
        f"{len(result.duplicados)} plasqlid duplicados ==="
    ]

    lines.append("")
    lines.append(
        f"Plasqlid que matchearon a más de una celda ({len(result.duplicados)}): "
        f"mismo texto repetido en más de una columna de la plantilla"
    )
    for d in result.duplicados:
        lines.append(f"  - plasqlid {d.plasqlid} -> celdas {list(d.celdas)}")
    if not result.duplicados:
        lines.append("  (ninguno)")

    lines.append("")
    lines.append(f"Celdas ambiguas ({len(result.ambiguos)}): mismo texto, varios plasqlid candidatos")
    for a in result.ambiguos:
        lines.append(f"  - {a.celda}: {a.texto!r} -> candidatos {list(a.candidatos)}")
    if not result.ambiguos:
        lines.append("  (ninguna)")

    lines.append("")
    lines.append(f"Celdas huérfanas ({len(result.huerfanos)}): texto sin match en PLACONSU.PLASQLDESC")
    for o in result.huerfanos:
        lines.append(f"  - {o.celda}: {o.texto!r}")
    if not result.huerfanos:
        lines.append("  (ninguna)")

    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def main(argv: Sequence[str] | None = None) -> None:
    from src.db.catalog_repository import fetch_catalog
    from src.db.connection import get_connection

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", required=True, help="ruta al .xlsx de la plantilla")
    parser.add_argument("--hoja", default=DEFAULT_HOJA)
    parser.add_argument("--institucion", required=True)
    parser.add_argument("--out-csv", default="indicator_map_propuesto.csv")
    parser.add_argument("--out-report", default="indicator_map_revision.txt")
    args = parser.parse_args(argv)

    conn = get_connection()
    try:
        catalog = fetch_catalog(conn)
    finally:
        conn.close()

    result = build_proposal_from_workbook(args.template, args.hoja, catalog)
    write_indicator_map_csv(result, args.institucion, args.hoja, args.out_csv)
    write_review_report(result, args.out_report)

    print(f"confirmados={len(result.confirmados)} ambiguos={len(result.ambiguos)} huerfanos={len(result.huerfanos)}")
    print(f"CSV propuesto: {args.out_csv}")
    print(f"Reporte de revisión: {args.out_report}")


if __name__ == "__main__":
    main()
