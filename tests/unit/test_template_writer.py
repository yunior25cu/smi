from pathlib import Path

import openpyxl
import pytest

from src.excel.template_writer import (
    TemplateWriterError,
    find_leftover_placeholder_cells,
    find_ordenes_ticket_columns,
    write_values,
)

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "sample_workbook.xlsx"
HOJA = "4315 Utilizacion "  # espacio final a propósito, igual que el archivo real


def test_fixture_exists():
    assert FIXTURE.exists(), "corré scripts/build_test_fixture.py para generarlo"


def test_writes_the_ordenes_ticket_cell(tmp_path):
    output = tmp_path / "out.xlsx"
    write_values(FIXTURE, output, HOJA, {"H13": 42})

    wb = openpyxl.load_workbook(output)
    sheet = wb[HOJA]
    assert sheet["H13"].value == 42


def test_precio_cell_is_never_touched(tmp_path):
    """La columna PRECIO (contigua a ORDENES/TICKET) no está en `values`
    nunca: tiene que quedar exactamente igual, valor y formato, antes y
    después de escribir."""
    output = tmp_path / "out.xlsx"
    write_values(FIXTURE, output, HOJA, {"H13": 42})

    wb = openpyxl.load_workbook(output)
    sheet = wb[HOJA]
    assert sheet["G13"].value == 100.5
    assert sheet["G13"].number_format == "0.00"


def test_does_not_modify_the_original_template(tmp_path):
    output = tmp_path / "out.xlsx"
    original_bytes = FIXTURE.read_bytes()

    write_values(FIXTURE, output, HOJA, {"H13": 42})

    assert FIXTURE.read_bytes() == original_bytes


def test_applies_default_integer_number_format_to_written_cells(tmp_path):
    output = tmp_path / "out.xlsx"
    write_values(FIXTURE, output, HOJA, {"H13": 42})

    wb = openpyxl.load_workbook(output)
    assert wb[HOJA]["H13"].number_format == "#,##0"


def test_number_format_can_be_overridden(tmp_path):
    output = tmp_path / "out.xlsx"
    write_values(FIXTURE, output, HOJA, {"H13": 42.5}, number_format="0.00")

    wb = openpyxl.load_workbook(output)
    assert wb[HOJA]["H13"].number_format == "0.00"


def test_unmapped_ordenes_ticket_cells_keep_their_placeholder_text(tmp_path):
    output = tmp_path / "out.xlsx"
    write_values(FIXTURE, output, HOJA, {"H13": 42})  # H14 no se mapea a propósito

    wb = openpyxl.load_workbook(output)
    sheet = wb[HOJA]
    assert sheet["H14"].value == "TEXTO PLACEHOLDER B"


def test_merged_vertical_concept_cells_are_preserved(tmp_path):
    output = tmp_path / "out.xlsx"
    write_values(FIXTURE, output, HOJA, {"H15": 1, "H16": 2})

    wb = openpyxl.load_workbook(output)
    sheet = wb[HOJA]
    assert "A15:A16" in [str(r) for r in sheet.merged_cells.ranges]


def test_sheet_protection_is_preserved_without_needing_the_password(tmp_path):
    output = tmp_path / "out.xlsx"
    write_values(FIXTURE, output, HOJA, {"H13": 42})

    wb = openpyxl.load_workbook(output)
    assert wb[HOJA].protection.sheet is True


def test_none_value_is_written_as_empty_cell(tmp_path):
    output = tmp_path / "out.xlsx"
    write_values(FIXTURE, output, HOJA, {"H13": None})

    wb = openpyxl.load_workbook(output)
    assert wb[HOJA]["H13"].value is None


def test_missing_template_raises_clear_error(tmp_path):
    output = tmp_path / "out.xlsx"
    with pytest.raises(TemplateWriterError, match="no existe"):
        write_values(tmp_path / "no_existe.xlsx", output, HOJA, {"H13": 1})


def test_missing_sheet_raises_clear_error(tmp_path):
    output = tmp_path / "out.xlsx"
    with pytest.raises(TemplateWriterError, match="HOJA_INEXISTENTE"):
        write_values(FIXTURE, output, "HOJA_INEXISTENTE", {"H13": 1})


def test_rerunning_for_the_same_output_overwrites_it(tmp_path):
    output = tmp_path / "out.xlsx"
    write_values(FIXTURE, output, HOJA, {"H13": 1})
    write_values(FIXTURE, output, HOJA, {"H13": 2})

    wb = openpyxl.load_workbook(output)
    assert wb[HOJA]["H13"].value == 2


# --- verificación post-escritura (celdas ORDENES/TICKET sin mapear) -------


def test_write_result_reports_leftover_placeholder_cells(tmp_path):
    output = tmp_path / "out.xlsx"
    # Solo se mapea H13: H14, J13, H15, H16 quedan con su texto original.
    result = write_values(FIXTURE, output, HOJA, {"H13": 42})

    leftover_cells = {celda for celda, _texto in result.celdas_con_texto_residual}
    assert leftover_cells == {"H14", "J13", "H15", "H16"}


def test_write_result_has_no_leftovers_when_everything_is_mapped(tmp_path):
    output = tmp_path / "out.xlsx"
    result = write_values(
        FIXTURE, output, HOJA, {"H13": 1, "H14": 2, "J13": 3, "H15": 4, "H16": 5}
    )

    assert result.celdas_con_texto_residual == ()


def test_find_ordenes_ticket_columns_detects_h_and_j_only():
    wb = openpyxl.load_workbook(FIXTURE)
    header_row, cols = find_ordenes_ticket_columns(wb[HOJA])

    letters = [wb[HOJA].cell(1, c).column_letter for c in cols]
    assert header_row == 12
    assert letters == ["H", "J"]


def test_find_leftover_placeholder_cells_does_not_depend_on_column_a():
    """Fila 16 tiene columna A vacía (parte del combinado A15:A16), pero
    igual tiene que aparecer: no hay que perder filas dentro de un
    combinado vertical."""
    wb = openpyxl.load_workbook(FIXTURE)
    leftovers = dict(find_leftover_placeholder_cells(wb[HOJA]))

    assert "H16" in leftovers
    assert leftovers["H16"] == "TEXTO PLACEHOLDER C2"
