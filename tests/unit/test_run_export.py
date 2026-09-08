from pathlib import Path

import openpyxl
import pytest

from src.core.indicator_map import IndicatorMapEntry
from src.core.run_export import run_export
from src.excel.workbook_paths import InstitutionConfig

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "sample_workbook.xlsx"
HOJA = "4315 Utilizacion "


class FakeCursor:
    def __init__(self, rows):
        self._rows = rows
        self.last_query = None

    def execute(self, query):
        self.last_query = query

    def fetchall(self):
        return self._rows


class FakeConnection:
    """Conexión DB-API 2.0 mockeada: ignora el SQL y devuelve filas fijas,
    como pide la Fase 4 para probar el orquestador sin pegarle a una base
    real."""

    def __init__(self, rows):
        self._rows = rows

    def cursor(self):
        return FakeCursor(self._rows)


def catalog_row(plasqlid, nombre="n", descripcion="d", sentencia="select 1 from t"):
    return (float(plasqlid), nombre, descripcion, sentencia)


@pytest.fixture
def institutions():
    return [InstitutionConfig(id="SMI", nombre="SMI", plantilla=FIXTURE, hoja=HOJA)]


def test_full_pipeline_writes_values_and_detects_all_discrepancies(tmp_path, institutions):
    # Catalogo real (via catalog_conn): indicadores 1, 2 y 3.
    catalog_conn = FakeConnection(
        [
            catalog_row(1, descripcion="Indicador uno"),
            catalog_row(2, descripcion="Indicador dos"),
            catalog_row(3, descripcion="Indicador tres, sin mapear"),
        ]
    )
    # Mapa: 1->H13 y 2->J13 (validos), 99->H99 (huerfano: no esta en el catalogo).
    indicator_map = [
        IndicatorMapEntry(plasqlid=1, institucion_id="SMI", hoja=HOJA, celda="H13"),
        IndicatorMapEntry(plasqlid=2, institucion_id="SMI", hoja=HOJA, celda="J13"),
        IndicatorMapEntry(plasqlid=99, institucion_id="SMI", hoja=HOJA, celda="H99"),
    ]
    # Resultado de la query consolidada (via dw_conn): 1 tiene valor, 2 da NULL.
    dw_conn = FakeConnection([(1, 12345.67), (2, None)])

    result = run_export(
        "SMI",
        2026,
        8,
        institutions=institutions,
        indicator_map=indicator_map,
        catalog_conn=catalog_conn,
        dw_conn=dw_conn,
        output_dir=tmp_path,
    )

    assert result.output_path == tmp_path / "SMI_2026_08.xlsx"
    assert result.output_path.exists()

    wb = openpyxl.load_workbook(result.output_path)
    sheet = wb[HOJA]
    assert sheet["H13"].value == 12345.67
    assert sheet["J13"].value is None  # NULL: se escribe como celda vacia
    # La columna PRECIO nunca se toca, aunque su ORDENES/TICKET si se escriba.
    assert sheet["G13"].value == 100.5
    # H14 nunca se mapeo: sigue con su texto original.
    assert sheet["H14"].value == "TEXTO PLACEHOLDER B"

    log = result.control_log
    assert [c.plasqlid for c in log.indicadores_sin_celda] == [3]
    assert [m.plasqlid for m in log.celdas_huerfanas] == [99]
    assert [m.plasqlid for m in log.indicadores_nulos] == [2]
    # H14, H15, H16 quedaron sin mapear (solo se mapearon H13 y J13).
    assert {c for c, _ in log.celdas_con_texto_residual} == {"H14", "H15", "H16"}

    assert result.control_log_path.exists()
    assert "Indicador tres, sin mapear" in result.control_log_path.read_text(encoding="utf-8")


def test_only_entries_for_the_target_institution_are_used(tmp_path, institutions):
    catalog_conn = FakeConnection([catalog_row(1)])
    indicator_map = [
        IndicatorMapEntry(plasqlid=1, institucion_id="SMI", hoja=HOJA, celda="H13"),
        IndicatorMapEntry(plasqlid=1, institucion_id="OTRA", hoja=HOJA, celda="H99"),
    ]
    dw_conn = FakeConnection([(1, 5.0)])

    result = run_export(
        "SMI",
        2026,
        8,
        institutions=institutions,
        indicator_map=indicator_map,
        catalog_conn=catalog_conn,
        dw_conn=dw_conn,
        output_dir=tmp_path,
    )

    wb = openpyxl.load_workbook(result.output_path)
    assert wb[HOJA]["H13"].value == 5.0
    assert result.control_log.indicadores_sin_celda == ()  # el 1 esta mapeado para SMI


def test_duplicate_plasqlid_mapped_to_same_cell_does_not_break_query_building(tmp_path, institutions):
    # Fase 1: dos filas identicas (mismo plasqlid+celda) son validas y se
    # toleran en indicator_map.csv. run_export tiene que deduplicar antes de
    # pasarle el catalogo a query_builder (que rechaza IDs repetidos).
    catalog_conn = FakeConnection([catalog_row(1)])
    indicator_map = [
        IndicatorMapEntry(plasqlid=1, institucion_id="SMI", hoja=HOJA, celda="H13"),
        IndicatorMapEntry(plasqlid=1, institucion_id="SMI", hoja=HOJA, celda="H13"),
    ]
    dw_conn = FakeConnection([(1, 7.0)])

    result = run_export(
        "SMI",
        2026,
        8,
        institutions=institutions,
        indicator_map=indicator_map,
        catalog_conn=catalog_conn,
        dw_conn=dw_conn,
        output_dir=tmp_path,
    )

    wb = openpyxl.load_workbook(result.output_path)
    assert wb[HOJA]["H13"].value == 7.0


def test_no_mapped_entries_still_produces_output_and_full_unmapped_log(tmp_path, institutions):
    catalog_conn = FakeConnection([catalog_row(1)])
    dw_conn = FakeConnection([])

    result = run_export(
        "SMI",
        2026,
        8,
        institutions=institutions,
        indicator_map=[],
        catalog_conn=catalog_conn,
        dw_conn=dw_conn,
        output_dir=tmp_path,
    )

    assert result.output_path.exists()
    assert [c.plasqlid for c in result.control_log.indicadores_sin_celda] == [1]
    # Ningun H/J se mapeo: las 5 celdas ORDENES/TICKET del fixture quedan
    # todas como texto residual.
    assert len(result.control_log.celdas_con_texto_residual) == 5
