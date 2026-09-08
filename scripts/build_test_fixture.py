"""Genera tests/fixtures/sample_workbook.xlsx: un workbook mínimo sintético
para los tests de template_writer.py y tools/build_indicator_map.py, con la
misma forma que la plantilla real (revisión 2 del prompt: hoja
"4315 Utilizacion " -con espacio final-, protegida, con pares de columnas
PRECIO/ORDENES-TICKET, celdas combinadas verticalmente, y number_format
propio en la columna PRECIO) pero sin ningún dato de producción.

Se corre una sola vez (o cuando haga falta ajustar el fixture); el archivo
resultante se versiona en git (es la única excepción a "no versionar .xlsx"
del .gitignore).

Uso:
    python scripts/build_test_fixture.py
"""

from pathlib import Path

import openpyxl

FIXTURE_PATH = Path(__file__).resolve().parent.parent / "tests" / "fixtures" / "sample_workbook.xlsx"
SHEET_NAME = "4315 Utilizacion "  # espacio final a propósito, como en el archivo real


def build() -> None:
    wb = openpyxl.Workbook()
    ref = wb.active
    ref.title = SHEET_NAME

    # Encabezados, igual que la plantilla real (fila 11 = "PRECIO N", fila 12
    # = "$" / "ORDENES/TICKET"). Dos pares de columnas alcanzan para probar
    # la detección de columnas y que no se toquen las que no corresponden.
    ref["A11"] = "CONCEPTO"
    ref["G11"] = "PRECIO 1"
    ref["I11"] = "PRECIO 2"
    ref["G12"] = "$"
    ref["H12"] = "ORDENES/TICKET"
    ref["I12"] = "$"
    ref["J12"] = "ORDENES/TICKET"

    # Fila 13: concepto A. PRECIO (G13) ya tiene un valor correcto que NUNCA
    # se debe tocar; ORDENES/TICKET (H13) tiene el placeholder típico (texto
    # con el PLASQLDESC del indicador) que hay que reemplazar por un número.
    ref["A13"] = "CONCEPTO A"
    ref["G13"] = 100.5
    ref["G13"].number_format = "0.00"
    ref["H13"] = "TEXTO PLACEHOLDER A"

    # Fila 14: concepto B, a propósito sin mapear en los tests -sirve para
    # probar que (a) la celda ORDENES/TICKET no mapeada queda con su texto
    # intacto, y (b) la verificación post-escritura la detecta como
    # "residual".
    ref["A14"] = "CONCEPTO B (no mapeado en los tests)"
    ref["G14"] = 200.0
    ref["G14"].number_format = "0.00"
    ref["H14"] = "TEXTO PLACEHOLDER B"

    # Filas 15-16: concepto C combinado verticalmente en columna A (como en
    # el archivo real: el texto del concepto vive solo en la primera celda
    # del combinado, pero cada fila tiene sus propios valores PRECIO/
    # ORDENES-TICKET). Prueba que la detección de placeholders no se pierda
    # filas por depender de la columna A.
    ref["A15"] = "CONCEPTO C (combinado verticalmente)"
    ref.merge_cells("A15:A16")
    ref["G15"] = 10.0
    ref["G15"].number_format = "0.00"
    ref["H15"] = "TEXTO PLACEHOLDER C1"
    ref["G16"] = 20.0
    ref["G16"].number_format = "0.00"
    ref["H16"] = "TEXTO PLACEHOLDER C2"

    # Columna PRECIO 2 (I/J): un valor de ejemplo para probar que escribir en
    # H no afecta a I/J y viceversa.
    ref["I13"] = 300.0
    ref["I13"].number_format = "0.00"
    ref["J13"] = "TEXTO PLACEHOLDER A (precio 2)"

    # Protección de hoja con contraseña real, como la plantilla real: hay
    # que verificar (Fase 0/3) que openpyxl puede escribir igual sin ella.
    ref.protection.sheet = True
    ref.protection.set_password("fixture-password")

    FIXTURE_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(FIXTURE_PATH)
    print(f"Generado: {FIXTURE_PATH}")


if __name__ == "__main__":
    build()
